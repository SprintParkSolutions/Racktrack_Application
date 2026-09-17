"""Build the Word version of the help-bot feedback form.

Mirrors help-bot-feedback-form.html question for question: same wording,
same help text, same worked examples. Answer areas are single-cell tables so
the form stays fillable in Word and still prints with a visible box.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x0E, 0x11, 0x14)
INK2 = RGBColor(0x33, 0x3B, 0x45)
MUTED = RGBColor(0x6E, 0x77, 0x83)
FAINT = RGBColor(0x98, 0xA1, 0xAD)
ACCENT = RGBColor(0x2E, 0x5F, 0x8F)
LINE = "D3D9E2"
SUNK = "E4E8EE"

BODY, MONO = "Calibri", "Consolas"


def shade(el, fill):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), fill)
    el.append(s)


def cell_borders(cell, color=LINE, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)
        borders.append(e)
    tcPr.append(borders)


def left_rule(par, color=LINE, sz=6, space=8):
    pPr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    bd.append(e)
    pPr.append(bd)


def bottom_rule(par, color=LINE, sz=6, space=6):
    pPr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), str(space))
    e.set(qn("w:color"), color)
    bd.append(e)
    pPr.append(bd)


def fix_width(table, pct=5000, cols=None):
    """Word autofits to content by default, which collapses an empty answer box
    to a sliver. Pin the table to a percentage of the text column and lay the
    grid out fixed.
    """
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(pct))
    w.set(qn("w:type"), "pct")
    tblPr.append(w)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    table.autofit = False
    if cols:
        grid = table._tbl.find(qn("w:tblGrid"))
        for gc, twips in zip(grid.findall(qn("w:gridCol")), cols):
            gc.set(qn("w:w"), str(twips))


def para(
    doc,
    text="",
    size=10.5,
    color=INK2,
    bold=False,
    italic=False,
    font=BODY,
    before=0,
    after=4,
    spacing=1.15,
    caps_track=False,
):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if text:
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.italic = italic
        if caps_track:
            rPr = r._element.get_or_add_rPr()
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), "30")  # tracked-out label, as on the page
            rPr.append(sp)
    return p


def answer_box(doc, lines=4, shaded=True):
    """A single-cell table people type into; prints as an empty ruled box."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    fix_width(t)
    c = t.rows[0].cells[0]
    cell_borders(c)
    if shaded:
        shade(c._tc.get_or_add_tcPr(), SUNK)
    c.text = ""
    for i in range(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.4
        r = p.add_run("")
        r.font.size = Pt(10.5)
        r.font.name = BODY
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def example(doc, label, body=None, items=None):
    p = para(doc, label, size=7.5, color=FAINT, font=MONO, before=2, after=3, caps_track=True)
    left_rule(p)
    if body:
        e = para(doc, body, size=9.5, color=MUTED, italic=True, after=6, spacing=1.25)
        left_rule(e)
    for i, it in enumerate(items or []):
        e = para(
            doc,
            f"{i + 1}.  {it}",
            size=9.5,
            color=MUTED,
            italic=True,
            after=(6 if i == len(items) - 1 else 1),
            spacing=1.2,
        )
        left_rule(e)


def question(doc, num, title, help_text=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    n = p.add_run(f"{num}   ")
    n.font.name = MONO
    n.font.size = Pt(9)
    n.font.color.rgb = ACCENT
    t = p.add_run(title)
    t.font.name = BODY
    t.font.size = Pt(12)
    t.font.color.rgb = INK
    t.bold = True
    if help_text:
        para(doc, help_text, size=9.5, color=MUTED, after=7)


def section(doc, title, note=None):
    p = para(doc, title, size=13, color=INK, bold=True, before=20, after=3)
    bottom_rule(p)
    if note:
        para(doc, note, size=9.5, color=MUTED, after=2)


# ----------------------------------------------------------------------------
doc = Document()
st = doc.styles["Normal"]
st.font.name = BODY
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)

s = doc.sections[0]
s.top_margin = s.bottom_margin = Inches(0.8)
s.left_margin = s.right_margin = Inches(0.9)

para(doc, "PRODUCT RESEARCH", size=7.5, color=FAINT, font=MONO, after=4, caps_track=True)

h = doc.add_paragraph()
h.paragraph_format.space_after = Pt(6)
r = h.add_run("If RackTrack had a help bot, what would you ask it?")
r.font.name = BODY
r.font.size = Pt(21)
r.font.color.rgb = INK
r.bold = True

para(
    doc,
    "RackTrack has no help bot today. Before we build one, tell us what you'd want "
    "from it — rough notes beat polished sentences.",
    size=11,
    color=INK2,
    after=10,
)

intro = para(
    doc,
    "Five questions, about three minutes. None of them are required. "
    "Fill this in, save it, and send it back to us — or print it and write on it.",
    size=9.5,
    color=MUTED,
    after=4,
)
bottom_rule(intro)

# -------------------------------------------------- 1
section(doc, "Who you are")
question(doc, "01", "Your name")
answer_box(doc, lines=1)

# -------------------------------------------------- 2
section(doc, "The main question", "The part we care about most. Take your time over these two.")

question(
    doc,
    "02",
    "If there were a support bot inside RackTrack, what problem would you expect it to solve for you?",
    "The thing you currently ask a colleague about, hunt through a document for, or give up on.",
)
example(
    doc,
    "EXAMPLE ANSWER",
    "“Half my time on site goes on working out whether the scan is wrong or the rack "
    "actually changed. I want to point at a result and ask why it thinks there's a switch "
    "in U12, and get a straight answer instead of re-scanning three times.”",
)
answer_box(doc, lines=6)

question(
    doc,
    "03",
    "Write five questions you'd actually type into it.",
    "As you would actually type them, mid-job. Don't clean them up.",
)
example(
    doc,
    "EXAMPLE ANSWERS",
    items=[
        "why is this scan still processing",
        "what was in rack 4 last month",
        "how do I add Ravi to the Bangalore site",
        "any free ports left on sw-03",
        "who changed the firmware last week",
    ],
)
t = doc.add_table(rows=5, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
fix_width(t, cols=(460, 9140))  # ~0.32in number column, rest for the answer
for i, row in enumerate(t.rows):
    numc, ansc = row.cells
    numc.width = Inches(0.32)
    ansc.width = Inches(6.0)
    cell_borders(numc, color="FFFFFF", sz=2)
    cell_borders(ansc)
    shade(ansc._tc.get_or_add_tcPr(), SUNK)
    np_ = numc.paragraphs[0]
    np_.paragraph_format.space_after = Pt(0)
    nr = np_.add_run(str(i + 1))
    nr.font.name = MONO
    nr.font.size = Pt(9)
    nr.font.color.rgb = FAINT
    ap = ansc.paragraphs[0]
    ap.paragraph_format.space_after = Pt(0)
    ap.paragraph_format.line_spacing = 1.3
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# -------------------------------------------------- 3
section(
    doc,
    "Where you actually get stuck",
    "One concrete example tells us more than ten feature requests.",
)

question(
    doc,
    "04",
    "Think of the last time RackTrack left you stuck or irritated. What were you trying to do?",
    "Where you were, what you expected to happen, what happened instead. Be blunt.",
)
example(
    doc,
    "EXAMPLE ANSWER",
    "“Tuesday, DC2 cold aisle. Trying to finish a multi-rack scan on the phone with "
    "gloves on, and it kept losing the site every time I came back from the camera. Did "
    "rack 3 twice, gave up, wrote it on paper and typed it in at my desk that evening.”",
)
answer_box(doc, lines=6)

question(
    doc,
    "05",
    "Which parts of RackTrack do you end up with questions about?",
    "Tick as many as apply. If something isn't listed, or you've never used it, note it in the box.",
)
example(
    doc,
    "EXAMPLE ANSWER",
    "“Ticked scan results and ports. Never opened Marketplace. If I could tick one of "
    "them twice it would be device detection — that's where the time goes.”",
)

OPTIONS = [
    "Scanning a rack with the camera",
    "Scan results & device detection",
    "Multi-rack scans",
    "Ports & available ports",
    "Connections & integrations",
    "Network view / live discovery",
    "Firmware checks",
    "Rack topology",
    "Switch info & specs",
    "Marketplace / buying SFPs",
    "Scan history",
    "Accounts, invites & permissions",
    "Organisation admin",
    "Signing in / installing / updates",
]
ct = doc.add_table(rows=7, cols=2)
ct.alignment = WD_TABLE_ALIGNMENT.LEFT
fix_width(ct, cols=(4800, 4800))
for i, opt in enumerate(OPTIONS):
    cell = ct.rows[i // 2].cells[i % 2]
    cell_borders(cell, color="FFFFFF", sz=2)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    box = p.add_run("☐  ")
    box.font.name = "Segoe UI Symbol"
    box.font.size = Pt(12)
    box.font.color.rgb = MUTED
    lab = p.add_run(opt)
    lab.font.name = BODY
    lab.font.size = Pt(10)
    lab.font.color.rgb = INK2
para(
    doc,
    "Anything else — or which of these you've never used:",
    size=9.5,
    color=MUTED,
    before=6,
    after=3,
)
answer_box(doc, lines=2)

# -------------------------------------------------- closing
p = para(
    doc, "SENDING IT BACK", size=7.5, color=ACCENT, font=MONO, before=18, after=3, caps_track=True
)
para(
    doc,
    "Save this file and send it back to us, or print it and write on it — "
    "whichever is easier. If you'd rather answer in a browser, the same five questions "
    "are on the RackTrack feedback page.",
    size=10,
    color=INK2,
    after=10,
)

f = para(doc, "RackTrack — help bot discovery", size=8.5, color=FAINT, before=6, after=0)

doc.save("/Volumes/Racktrack/dark_mobile/docs/support-bot/help-bot-feedback-form.docx")
print("written")
